import csv
import io
import json
from dataclasses import dataclass
from pathlib import PurePath

from app.textutil import estimate_tokens, split_sentences


@dataclass
class Section:
    heading: str
    text: str


@dataclass
class Chunk:
    ord: int
    heading: str
    content: str


def _extract_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return [Section(f"Page {i}", page.extract_text() or "") for i, page in enumerate(reader.pages, 1)]


def _extract_docx(data: bytes) -> list[Section]:
    from docx import Document

    document = Document(io.BytesIO(data))
    sections: list[Section] = []
    heading, buffer = "", []
    for paragraph in document.paragraphs:
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.lower().startswith("heading") and paragraph.text.strip():
            if buffer:
                sections.append(Section(heading, "\n".join(buffer)))
            heading, buffer = paragraph.text.strip(), []
        elif paragraph.text.strip():
            buffer.append(paragraph.text.strip())
    if buffer:
        sections.append(Section(heading, "\n".join(buffer)))
    for index, table in enumerate(document.tables, 1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        sections.append(Section(f"Table {index}", "\n".join(rows)))
    return sections


def _extract_markdown(text: str) -> list[Section]:
    sections: list[Section] = []
    path: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        if any(line.strip() for line in buffer):
            sections.append(Section(" > ".join(path), "\n".join(buffer).strip()))
        buffer.clear()

    for line in text.splitlines():
        if line.startswith("#"):
            flush()
            level = len(line) - len(line.lstrip("#"))
            path[:] = path[: level - 1] + [line.lstrip("#").strip()]
        else:
            buffer.append(line)
    flush()
    return sections


def _extract_csv(text: str, rows_per_section: int = 25) -> list[Section]:
    reader = list(csv.reader(io.StringIO(text)))
    if not reader:
        return []
    header, rows = reader[0], reader[1:]
    sections = []
    for start in range(0, len(rows), rows_per_section):
        lines = ["; ".join(f"{h}: {v}" for h, v in zip(header, row)) + "." for row in rows[start : start + rows_per_section]]
        sections.append(Section(f"Rows {start + 1}-{start + len(lines)}", "\n".join(lines)))
    return sections


def _extract_json(text: str) -> list[Section]:
    data = json.loads(text)
    if isinstance(data, dict):
        return [Section(str(key), json.dumps(value, indent=1, ensure_ascii=False)) for key, value in data.items()]
    if isinstance(data, list):
        return [Section(f"Item {i}", json.dumps(item, indent=1, ensure_ascii=False)) for i, item in enumerate(data, 1)]
    return [Section("", str(data))]


def extract(filename: str, data: bytes) -> list[Section]:
    extension = PurePath(filename).suffix.lower()
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension == ".docx":
        return _extract_docx(data)
    text = data.decode("utf-8")
    if extension == ".md":
        return _extract_markdown(text)
    if extension == ".csv":
        return _extract_csv(text)
    if extension == ".json":
        return _extract_json(text)
    return [Section(PurePath(filename).stem, text)]


def _split_long(sentence: str, max_tokens: int) -> list[str]:
    words = sentence.split()
    step = max(1, int(max_tokens / 1.3))
    return [" ".join(words[i : i + step]) for i in range(0, len(words), step)]


def chunk_sections(sections: list[Section], max_tokens: int = 350, overlap_sentences: int = 1) -> list[Chunk]:
    """Structure-aware chunking: never crosses section boundaries, sentence-aligned, with sentence overlap."""
    chunks: list[Chunk] = []
    for section in sections:
        sentences: list[str] = []
        for sentence in split_sentences(section.text):
            sentences.extend(_split_long(sentence, max_tokens) if estimate_tokens(sentence) > max_tokens else [sentence])
        buffer: list[str] = []
        buffer_tokens = 0
        for sentence in sentences:
            tokens = estimate_tokens(sentence)
            if buffer and buffer_tokens + tokens > max_tokens:
                chunks.append(Chunk(len(chunks), section.heading, " ".join(buffer)))
                buffer = buffer[-overlap_sentences:] if overlap_sentences else []
                buffer_tokens = sum(estimate_tokens(s) for s in buffer)
            buffer.append(sentence)
            buffer_tokens += tokens
        if buffer:
            chunks.append(Chunk(len(chunks), section.heading, " ".join(buffer)))
    return chunks
